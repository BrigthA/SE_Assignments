import json
import re
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime

#read data from excelsheet(Excel file named contacts.xlsx with columns: Name, Phone (with country code, e.g. +91xxxxxxxxxx) and send message to whatsapp web and extract the data from Whatsapp web and save it in excel sheet and json

with sync_playwright() as p:
    profile_path = Path.home() / "whatsapp_playwright_profile"
    context = p.chromium.launch_persistent_context(user_data_dir=str(profile_path),headless=False,)
    page = context.pages[0] if context.pages else context.new_page()
    page.goto("https://web.whatsapp.com/")  # Navigates to WhatsApp Web
    search_box = page.get_by_role("textbox", name="Search or start a new chat")
    search_box.wait_for(timeout=120000)  # Waits for login and the chat list to be ready

    # Load the Excel workbook and select the active sheet
    workbook = load_workbook(filename="contacts.xlsx")
    sheet = workbook.active
    message_records = []

    # Iterate through each row in the Excel sheet (starting from row 2 to skip headers)
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = row[0] if len(row) > 0 else None
        phone = row[1] if len(row) > 1 else None
        message_template = row[2] if len(row) > 2 else None
        if all(value is None or not str(value).strip() for value in row):
            continue

        try:
            if name is None or phone is None or not str(phone).strip():
                raise ValueError("Name and phone are required")

            message = (
                str(message_template).format(name=name, phone=phone)
                if message_template is not None and str(message_template).strip()
                else f"Hello {name}, this is a test message from Playwright Bot."
            )

            phone_digits = re.sub(r"\D", "", str(phone))
            if not phone_digits:
                raise ValueError("Phone number contains no digits")

            page.goto(f"https://web.whatsapp.com/send?phone={phone_digits}")
            message_box = page.get_by_role("textbox", name="Type a message")
            message_box.wait_for(timeout=30000)
            message_box.fill(message)
            page.keyboard.press("Enter")
            page.wait_for_timeout(5000)
            print(f"Message sent to {name} ({phone})")

            message_containers = page.locator('[data-testid="msg-container"]')
            message_count = message_containers.count()
            first_message = max(0, message_count - 3)
            for message_index in range(first_message, message_count):
                message_element = message_containers.nth(message_index)
                text_parts = message_element.locator('[data-testid="selectable-text"]').all_inner_texts()
                message_text = "\n".join(part.strip() for part in text_parts if part.strip())
                if not message_text:
                    message_text = message_element.inner_text().strip()

                message_records.append({
                    "name": str(name),
                    "phone": str(phone),
                    "message_number": message_index - first_message + 1,
                    "message": message_text,
                    "status": "sent",
                    "error": "",
                })

            if message_count == 0:
                message_records.append({
                    "name": str(name),
                    "phone": str(phone),
                    "message_number": None,
                    "message": "",
                    "status": "sent; no messages extracted",
                    "error": "",
                })
        except Exception as error:
            print(f"Failed to process Excel row {row_number} ({name}, {phone}): {error}")
            message_records.append({
                "name": str(name) if name is not None else "",
                "phone": str(phone) if phone is not None else "",
                "message_number": None,
                "message": "",
                "status": "error",
                "error": str(error),
            })

    workbook.close()

    output_path = Path(__file__).parent
    output_workbook = Workbook()
    output_sheet = output_workbook.active
    output_sheet.title = "Report"
    output_sheet.append(["Name", "Phone", "Message Number", "Message", "Status", "Error"])
    for record in message_records:
        output_sheet.append([
            record["name"],
            record["phone"],
            record["message_number"],
            record["message"],
            record["status"],
            record["error"],
        ])
    date = datetime.now().strftime("%Y-%m-%d")
    output_workbook.save(output_path / f"whatsapp_report_{date}.xlsx")

    with open(output_path / f"whatsapp_report_{date}.json", "w", encoding="utf-8") as json_file:
        json.dump(message_records, json_file, ensure_ascii=False, indent=2)
